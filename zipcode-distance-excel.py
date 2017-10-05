from pyzipcode import ZipCodeDatabase
from geopy.distance import vincenty
from openpyxl import load_workbook
from string import ascii_uppercase
from os import listdir, path
import sys


def col2num(col):
    """
    :param col: the string version of the column number you want. i.e A -> 1, B -> 2, C -> 3 etc
    :return: the index of the column (one-indexed)
    """

    num = 0
    for c in col:
        if c in ascii_uppercase:
            num = num * len(ascii_uppercase) + (ord(c) - ord(ascii_uppercase[0])) + 1
    return num


if __name__ == "__main__":

    files_in_dir = listdir(path.abspath(path.dirname(__file__)))

    xl_files_in_dir = []

    for file in files_in_dir:
        if file.endswith(".xlsx"):
            xl_files_in_dir.append(file)

    if len(xl_files_in_dir) <= 0:
        print("There aren't any .xlsx files in this directory! Exiting.")
        exit(0)

    print("Which file would you like to modify?")

    for file_number, file_name in enumerate(xl_files_in_dir):
        print("[" + str(file_number) + "] - ", file_name)

    file_number_string = input("File Number: ")

    file_number = None

    try:
        file_number = int(file_number_string)
    except ValueError:
        print("[" + str(file_number_string) + "]", "Isn't a number. Exiting.")
        exit(0)

    selected_file_name = None

    try:
        selected_file_name = xl_files_in_dir[file_number]
    except IndexError:
        print("[" + str(file_number) + "]", "Doesn't correspond with a listed file number. Exiting.")

    print("You've selected [" + str(selected_file_name) + "]", "to edit.")

    wb = load_workbook(filename=selected_file_name)

    print("Which sheet would you like to modify?")

    for sheet_number, sheet_name in enumerate(wb.sheetnames):
        print("[" + str(file_number) + "] - ", sheet_name)

    sheet_name = 0
    zip_column = 'I'
    target_column = 'J'


    ws = wb[sheet_name]

    zipcode_database = ZipCodeDatabase()
    mgh_zip = '02114'
    mgh = zipcode_database[mgh_zip]
    mgh_coords = (mgh.latitude, mgh.longitude)

    for row in ws.iter_rows(min_row=1):

        for cell in row:

                if cell.column == zip_column:

                    if cell.value is not None:

                        lookup_zip = str(cell.value).zfill(5)  # there are 5 digits in a zipcode

                        try:
                            zipcode = zipcode_database[lookup_zip]

                            distance = vincenty(mgh_coords, (zipcode.latitude, zipcode.longitude)).miles

                            insertion_text = distance

                        except IndexError as e:
                            print("Couldn't find zipcode:",
                                  str(lookup_zip),
                                  "at cell: " + str(cell.column) + str(cell.row),
                                  "in Database")

                            insertion_text = "?"

                        target_column_index = col2num(target_column)
                        write_cell = ws.cell(row=cell.row, column=target_column_index)
                        write_cell.value = insertion_text

    wb.save(selected_file_name)




